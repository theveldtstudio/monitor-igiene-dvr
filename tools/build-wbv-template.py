"""
Build public/templates/wbv.xlsx — Foglio di campagna ufficiale Vibrazioni WBV.
Pattern: foglio unico "CI", ~200 misure.

Colonne (A..Q):
  A  REG | [n°]
  B  DATA
  C  DURATA
  D  MACCHINA/MEZZO
  E  TARGA
  F  POSIZIONE OPERATORE
  G  TRAZIONE
  H  UTENSILE/OPERAZIONE
  I  REGIME
  J  FASE LAVORATIVA
  K  aw X (m/s²)
  L  aw Y (m/s²)
  M  aw Z (m/s²)
  N  A(w)max [formula]
  O  ASSE [formula]
  P  NOTE
  Q  TEMPERATURA
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

OUT = Path(__file__).resolve().parent.parent / "public" / "templates" / "wbv.xlsx"

COL_WIDTHS = {
    "A": 6.0,
    "B": 12.0,
    "C": 12.0,
    "D": 22.0,
    "E": 12.0,
    "F": 16.0,
    "G": 12.0,
    "H": 16.0,
    "I": 10.0,
    "J": 16.0,
    "K": 9.0,
    "L": 9.0,
    "M": 9.0,
    "N": 10.0,
    "O": 8.0,
    "P": 20.0,
    "Q": 12.0,
}
NUM_COLS = 17
DATA_ROWS = 200
DATA_START_ROW = 5

THIN = Side(style="thin", color="000000")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")
FONT_HDR    = Font(name="Calibri", size=10, bold=True)
FONT_TITLE  = Font(name="Calibri", size=12, bold=True)
FONT_NORMAL = Font(name="Calibri", size=10)
FILL_TITLE  = PatternFill(fill_type="solid", fgColor="DDDDDD")
FILL_COLHDR = PatternFill(fill_type="solid", fgColor="EEEEEE")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "CI"

    for col_letter, w in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = w

    # --- Riga 1: header anagrafica ---
    ws.row_dimensions[1].height = 30.0
    ws.cell(1, 1, "Impresa:").font = FONT_HDR
    ws.cell(1, 1).alignment = CENTER
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
    ws.cell(1, 2).alignment = CENTER
    ws.cell(1, 4, "Data:").font = FONT_HDR
    ws.cell(1, 4).alignment = CENTER
    ws.cell(1, 5, "Cantiere:").font = FONT_HDR
    ws.cell(1, 5).alignment = CENTER
    ws.merge_cells(start_row=1, start_column=6, end_row=1, end_column=14)
    ws.cell(1, 6).alignment = CENTER
    ws.cell(1, 15).value = None  # spacer
    ws.cell(1, 16, "Progressiva:").font = FONT_HDR
    ws.cell(1, 16).alignment = CENTER
    ws.cell(1, 17, "Pag: 1").font = FONT_HDR
    ws.cell(1, 17).alignment = RIGHT

    # --- Riga 2: titolo ---
    ws.row_dimensions[2].height = 24.95
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NUM_COLS)
    c = ws.cell(2, 1, "CAMPIONAMENTO VIBRAZIONI CORPO INTERO (WBV) (Titolo VIII D.Lgs 81/08)")
    c.font = FONT_TITLE
    c.alignment = CENTER
    c.fill = FILL_TITLE

    # --- Riga 3: header colonne (con merge verticale per colonne singole e merge
    #     orizzontale per ACCELERAZIONE) ---
    ws.row_dimensions[3].height = 30.0
    ws.row_dimensions[4].height = 18.0

    # Colonne con span verticale r3:r4
    vert_span_cols = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 14, 15, 16, 17]
    vert_labels = {
        1:  "REG\n[n°]",
        2:  "DATA",
        3:  "DURATA",
        4:  "MACCHINA/\nMEZZO",
        5:  "TARGA",
        6:  "POSIZIONE\nOPERATORE",
        7:  "TRAZIONE",
        8:  "UTENSILE/\nOPERAZIONE",
        9:  "REGIME",
        10: "FASE\nLAVORATIVA",
        14: "A(w)max\n(m/s²)",
        15: "ASSE",
        16: "NOTE",
        17: "TEMPERATURA\n(°C)",
    }
    for col in vert_span_cols:
        ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
        c = ws.cell(3, col, vert_labels[col])
        c.font = FONT_HDR
        c.alignment = CENTER
        c.fill = FILL_COLHDR

    # ACCELERAZIONE aw (m/s²) → K:M r3 merge orizzontale
    ws.merge_cells(start_row=3, start_column=11, end_row=3, end_column=13)
    c = ws.cell(3, 11, "ACCELERAZIONE aw (m/s²)")
    c.font = FONT_HDR
    c.alignment = CENTER
    c.fill = FILL_COLHDR

    # Sub-header r4: X, Y, Z
    for col, lbl in [(11, "X"), (12, "Y"), (13, "Z")]:
        c = ws.cell(4, col, lbl)
        c.font = FONT_HDR
        c.alignment = CENTER
        c.fill = FILL_COLHDR

    # Bordi su r3:r4 per tutte le colonne
    for rr in (3, 4):
        for col in range(1, NUM_COLS + 1):
            ws.cell(rr, col).border = BORDER_ALL

    # --- Righe dati 5..204 ---
    last_data_row = DATA_START_ROW + DATA_ROWS - 1
    for i in range(DATA_ROWS):
        rr = DATA_START_ROW + i
        ws.row_dimensions[rr].height = 18.0
        for col in range(1, NUM_COLS + 1):
            ws.cell(rr, col).border = BORDER_ALL
            ws.cell(rr, col).alignment = CENTER
            ws.cell(rr, col).font = FONT_NORMAL
        # Numero misura in A
        ws.cell(rr, 1, i + 1).font = FONT_HDR
        # Formula N = A(w)max
        k, l, m = f"K{rr}", f"L{rr}", f"M{rr}"
        ws.cell(rr, 14).value = f'=IF(AND({k}="",{l}="",{m}=""),"",MAX({k},{l},{m}))'
        # Formula O = asse dominante
        n = f"N{rr}"
        ws.cell(rr, 15).value = (
            f'=IF({n}="","",IF(AND({k}>={l},{k}>={m}),"X",IF({l}>={m},"Y","Z")))'
        )

    # --- Footer ---
    r_foot = last_data_row + 1
    ws.row_dimensions[r_foot].height = 27.75
    ws.merge_cells(start_row=r_foot, start_column=1, end_row=r_foot, end_column=2)
    c = ws.cell(r_foot, 1, "Tecnico rilevatore")
    c.font = FONT_HDR
    c.alignment = LEFT
    ws.merge_cells(start_row=r_foot, start_column=3, end_row=r_foot, end_column=10)
    ws.cell(r_foot, 3).alignment = LEFT
    ws.merge_cells(start_row=r_foot, start_column=11, end_row=r_foot, end_column=NUM_COLS)
    c = ws.cell(r_foot, 11, "Strumentazione:")
    c.font = FONT_NORMAL
    c.alignment = LEFT

    # --- Page setup ---
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:4"

    from openpyxl.utils import get_column_letter
    last_col = get_column_letter(NUM_COLS)
    ws.print_area = f"A1:{last_col}{r_foot}"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
