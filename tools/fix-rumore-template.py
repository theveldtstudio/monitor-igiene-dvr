#!/usr/bin/env python3
"""
Fix strutturale public/templates/rumore.xlsx:
- Col B (DURATA) righe dati: numFmt '[hh]:mm:ss'
- Col F, G, H (Leq dBA, Leq dBC, Lpeak) righe dati: numFmt '0.0'
- Col A (REG), C, D, E, I (testo): numFmt 'General' esplicito per forzare xf separati

4 pagine x 8 misure:
  Pagina 1: righe 4-11
  Pagina 2: righe 19-26
  Pagina 3: righe 34-41
  Pagina 4: righe 49-56
"""
import shutil
import openpyxl

TEMPLATE_PATH = "public/templates/rumore.xlsx"
BACKUP_PATH   = "public/templates/rumore.xlsx.bak"

DATA_ROWS_PER_PAGE = [
    range(4,  12),   # pagina 1: righe 4-11
    range(19, 27),   # pagina 2: righe 19-26
    range(34, 42),   # pagina 3: righe 34-41
    range(49, 57),   # pagina 4: righe 49-56
]

NUM_FMT_DURATA = '[hh]:mm:ss'
NUM_FMT_NUM    = '0.0'
NUM_FMT_TEXT   = 'General'


def main():
    print(f"Backup: {TEMPLATE_PATH} -> {BACKUP_PATH}")
    shutil.copy(TEMPLATE_PATH, BACKUP_PATH)

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.active

    changed = 0
    for page_rows in DATA_ROWS_PER_PAGE:
        for r in page_rows:
            # B: DURATA — orario
            b = ws.cell(row=r, column=2)
            b.number_format = NUM_FMT_DURATA
            changed += 1

            # F, G, H: Leq dBA, Leq dBC, Lpeak — decimali
            for col in (6, 7, 8):
                ws.cell(row=r, column=col).number_format = NUM_FMT_NUM
                changed += 1

            # A, C, D, E, I: testo — esplicito per forzare xf separati
            for col in (1, 3, 4, 5, 9):
                ws.cell(row=r, column=col).number_format = NUM_FMT_TEXT
                changed += 1

    wb.save(TEMPLATE_PATH)
    print(f"Fix applicato: {changed} celle modificate.")
    print(f"Salvato {TEMPLATE_PATH}")


if __name__ == "__main__":
    main()
