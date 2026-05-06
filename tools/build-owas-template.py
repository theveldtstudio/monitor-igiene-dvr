"""
Build public/templates/owas.xlsx — Foglio di campagna ufficiale OWAS.
Pattern: foglio unico "Foglio1", 8 misure (da 9 → 8).

Layout:
  Riga 1  : header anagrafica
  Riga 2  : titolo "OWAS (Titolo VI D.Lgs 81/08)"
  Riga 3-4: header colonne (con sub-header CODICE OWAS)
  Righe 5-12: dati (max 8 misure)
  Riga 13 : separatore vuoto
  Riga 14 : footer tecnico rilevatore

Colonne (A..J):
  A  REG [n°]
  B  MANSIONE
  C  ATTIVITA'
  D  DURATA
  E  SCHIENA  ─┐
  F  BRACCIA   │ CODICE OWAS
  G  GAMBE     │
  H  CARICO    │
  I  CLASSE   ─┘
  J  NOTE
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

OUT = Path(__file__).resolve().parent.parent / "public" / "templates" / "owas.xlsx"

COL_WIDTHS = {
    "A": 6.0,
    "B": 22.0,
    "C": 25.0,
    "D": 12.0,
    "E": 10.0,
    "F": 10.0,
    "G": 10.0,
    "H": 10.0,
    "I": 10.0,
    "J": 30.0,
}
NUM_COLS = 10
DATA_ROWS = 8
DATA_START_ROW = 5

THIN = Side(style="thin", color="000000")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")
FONT_HDR    = Font(name="Calibri", size=10, bold=True)
FONT_TITLE  = Font(name="Calibri", size=12, bold=True)
FONT_NORMAL = Font(name="Calibri", size=10)
FILL_TITLE  = PatternFill(fill_type="solid", fgColor="DDDDDD")
FILL_COLHDR = PatternFill(fill_type="solid", fgColor="EEEEEE")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Foglio1"

    for col_letter, w in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = w

    # --- Riga 1: header anagrafica ---
    # Export schema scrive: B1=committente, C1="Data: gg/mm/aaaa", E1=cantiere
    ws.row_dimensions[1].height = 30.0
    ws.cell(1, 1, "Impresa:").font = FONT_HDR
    ws.cell(1, 1).alignment = CENTER
    ws.cell(1, 2).alignment = CENTER                        # committente → scritto da export
    ws.cell(1, 3).alignment = CENTER                        # data → scritto da export
    ws.cell(1, 4, "Cantiere:").font = FONT_HDR
    ws.cell(1, 4).alignment = CENTER
    ws.merge_cells(start_row=1, start_column=5, end_row=1, end_column=8)
    ws.cell(1, 5).alignment = CENTER                        # cantiere → scritto da export (E1)
    ws.cell(1, 9, "Pag: 1").font = FONT_HDR
    ws.cell(1, 9).alignment = RIGHT

    # --- Riga 2: titolo ---
    ws.row_dimensions[2].height = 24.95
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NUM_COLS)
    c = ws.cell(2, 1, "OWAS (Titolo VI D.Lgs 81/08)")
    c.font = FONT_TITLE
    c.alignment = CENTER
    c.fill = FILL_TITLE

    # --- Righe 3-4: header colonne ---
    ws.row_dimensions[3].height = 30.0
    ws.row_dimensions[4].height = 18.0

    # Colonne con span verticale r3:r4
    vert_span = {
        1: "REG\n[n°]",
        2: "MANSIONE",
        3: "ATTIVITA'",
        4: "DURATA",
        10: "NOTE",
    }
    for col, lbl in vert_span.items():
        ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
        c = ws.cell(3, col, lbl)
        c.font = FONT_HDR
        c.alignment = CENTER
        c.fill = FILL_COLHDR

    # CODICE OWAS → E:I r3 merge orizzontale
    ws.merge_cells(start_row=3, start_column=5, end_row=3, end_column=9)
    c = ws.cell(3, 5, "CODICE OWAS")
    c.font = FONT_HDR
    c.alignment = CENTER
    c.fill = FILL_COLHDR

    # Sub-header r4
    sub4 = {5: "SCHIENA", 6: "BRACCIA", 7: "GAMBE", 8: "CARICO", 9: "CLASSE"}
    for col, lbl in sub4.items():
        c = ws.cell(4, col, lbl)
        c.font = FONT_HDR
        c.alignment = CENTER
        c.fill = FILL_COLHDR

    # Bordi r3:r4
    for rr in (3, 4):
        for col in range(1, NUM_COLS + 1):
            ws.cell(rr, col).border = BORDER_ALL

    # --- Righe dati 5-12 (8 slot) ---
    for i in range(DATA_ROWS):
        rr = DATA_START_ROW + i
        ws.row_dimensions[rr].height = 30.0
        for col in range(1, NUM_COLS + 1):
            ws.cell(rr, col).border = BORDER_ALL
            ws.cell(rr, col).alignment = CENTER
            ws.cell(rr, col).font = FONT_NORMAL
        ws.cell(rr, 1, i + 1).font = FONT_HDR

    # --- Riga 13: separatore vuoto ---
    ws.row_dimensions[13].height = 15.0

    # --- Riga 14: footer tecnico ---
    ws.row_dimensions[14].height = 27.75
    ws.merge_cells(start_row=14, start_column=1, end_row=14, end_column=2)
    c = ws.cell(14, 1, "Tecnico rilevatore:")
    c.font = FONT_HDR
    c.alignment = LEFT
    ws.merge_cells(start_row=14, start_column=3, end_row=14, end_column=NUM_COLS)
    ws.cell(14, 3).alignment = LEFT

    # --- Page setup ---
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:J14"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
