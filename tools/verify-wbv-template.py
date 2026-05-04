import openpyxl

TEMPLATE_PATH = "public/templates/wbv.xlsx"
SHEET_NAME = "CI"

wb = openpyxl.load_workbook(TEMPLATE_PATH)
ws = wb[SHEET_NAME]

print("=== a) Merge cells in sheet 'CI' ===")
merges = sorted(str(mr) for mr in ws.merged_cells.ranges)
for m in merges:
    print(f"  {m}")
print(f"  Total: {len(merges)}")

print("\n=== b) Column N width ===")
col_n = ws.column_dimensions.get("N")
width = col_n.width if col_n else "(not set)"
print(f"  Column N width: {width}")

print("\n=== c) Formulas N5 and O5 ===")
n5 = ws["N5"].value
o5 = ws["O5"].value
print(f"  N5 = {repr(n5)}")
print(f"  O5 = {repr(o5)}")
if n5 and str(n5).startswith("="):
    print("  N5: formula OK")
else:
    print("  N5: WARNING — no formula found")
if o5 and str(o5).startswith("="):
    print("  O5: formula OK")
else:
    print("  O5: WARNING — no formula found")

print("\n=== d) Progressive numbers A5+ ===")
issues = []
for row in range(5, 15):
    val = ws.cell(row=row, column=1).value
    expected = row - 4
    status = "OK" if val == expected else f"WARN (got {repr(val)}, expected {expected})"
    print(f"  A{row} = {repr(val)}  {status}")
    if val != expected:
        issues.append(row)

print()
if issues:
    print(f"WARNING: rows with unexpected values: {issues}")
else:
    print("All A5:A14 progressive numbers intact.")
