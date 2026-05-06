"""
Build public/templates/biologico-sas.xlsx — Foglio di campagna ufficiale Biologico SAS.
4 pagine × 8 misure = 32 max.

Colonne (A..K):
  A  REG [n°]
  B  COD. FILTRO
  C  POSTAZIONE
  D  FASE / MANSIONE
  E  MACCHINE/IMPIANTI
  F  T. PRELIEVO
  G  VOLUME (L)
  H  CONTA 22°C (UFC/m³)
  I  CONTA 36°C (UFC/m³)
  J  MUFFE+LIEVITI (UFC/m³)
  K  NOTE

Nota: i campi H/I/J supportano la stringa "<X" per valori sotto soglia (LoQ).
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.page import PageMargins

OUT = Path(__file__).resolve().parent.parent / "public" / "templates" / "biologico-sas.xlsx"

COL_WIDTHS = {
    "A": 7.0,
    "B": 12.0,
    "C": 18.0,
    "D": 18.0,
    "E": 20.0,
    "F": 12.0,
    "G": 11.0,
    "H": 14.0,
    "I": 14.0,
    "J": 16.0,
    "K": 22.0,
}
ROW_H_HEADER = 30.0
ROW_H_TITLE  = 24.95
ROW_H_COLHDR = 36.0
ROW_H_DATA   = 36.0
ROW_H_FOOTER = 27.75
ROW_H_EMPTY  = 15.0

PAGE_ROWS = 13
DATA_ROWS_PER_PAGE = 8
NUM_COLS = 11
NUM_PAGES = 4

THIN = Side(style="thin", color="000000")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
FONT_HDR    = Font(name="Calibri", size=10, bold=True)
FONT_TITLE  = Font(name="Calibri", size=12, bold=True)
FONT_NORMAL = Font(name="Calibri", size=10)
FILL_TITLE  = PatternFill(fill_type="solid", fgColor="DDDDDD")
FILL_COLHDR = PatternFill(fill_type="solid", fgColor="EEEEEE")

HEADERS = [
    "REG\n[n°]",
    "COD.\nFILTRO",
    "POSTAZIONE",
    "FASE /\nMANSIONE",
    "MACCHINE /\nIMPIANTI",
    "T.\nPRELIEVO",
    "VOLUME\n(L)",
    "CONTA 22°C\n(UFC/m³)",
    "CONTA 36°C\n(UFC/m³)",
    "MUFFE+LIEVITI\n(UFC/m³)",
    "NOTE",
]


def build_page(ws, page_num: int, base_row: int, slot_start: int):
    r = base_row
    ws.row_dimensions[r].height = ROW_H_HEADER
    ws.cell(r, 1, "Impresa:").font = FONT_HDR
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(r, 2).alignment = CENTER
    ws.cell(r, 4, "Data:").font = FONT_HDR
    ws.cell(r, 5, "Cantiere:").font = FONT_HDR
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
    ws.cell(r, 6).alignment = CENTER
    ws.cell(r, 10, "Progressiva:").font = FONT_HDR
    ws.cell(r, 11, f"Pag: {page_num}").font = FONT_HDR
    ws.cell(r, 11).alignment = Alignment(horizontal="right", vertical="center")

    r2 = base_row + 1
    ws.row_dimensions[r2].height = ROW_H_TITLE
    ws.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=NUM_COLS)
    c = ws.cell(r2, 1, "MONITORAGGIO AGENTI BIOLOGICI — SAS (D.Lgs 81/08 — Titolo X)")
    c.font = FONT_TITLE
    c.alignment = CENTER
    c.fill = FILL_TITLE

    r3 = base_row + 2
    ws.row_dimensions[r3].height = ROW_H_COLHDR
    for col, txt in enumerate(HEADERS, start=1):
        c = ws.cell(r3, col, txt)
        c.font = FONT_HDR
        c.alignment = CENTER
        c.fill = FILL_COLHDR
        c.border = BORDER_ALL

    for i in range(DATA_ROWS_PER_PAGE):
        rr = base_row + 3 + i
        ws.row_dimensions[rr].height = ROW_H_DATA
        for col in range(1, NUM_COLS + 1):
            ws.cell(rr, col).border = BORDER_ALL
            ws.cell(rr, col).alignment = CENTER
            ws.cell(rr, col).font = FONT_NORMAL
        ws.cell(rr, 1, slot_start + i).font = FONT_HDR

    r_foot = base_row + 11
    ws.row_dimensions[r_foot].height = ROW_H_FOOTER
    ws.merge_cells(start_row=r_foot, start_column=1, end_row=r_foot, end_column=2)
    c = ws.cell(r_foot, 1, "Tecnico rilevatore")
    c.font = FONT_HDR
    c.alignment = LEFT
    ws.merge_cells(start_row=r_foot, start_column=3, end_row=r_foot, end_column=5)
    ws.cell(r_foot, 3).alignment = LEFT
    ws.merge_cells(start_row=r_foot, start_column=6, end_row=r_foot, end_column=NUM_COLS)
    c = ws.cell(r_foot, 6, "Strumentazione:")
    c.font = FONT_NORMAL
    c.alignment = LEFT

    r_sep = base_row + 12
    if r_sep <= base_row + PAGE_ROWS - 1:
        ws.row_dimensions[r_sep].height = ROW_H_EMPTY


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Foglio1"

    for col_letter, w in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = w

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True

    breaks = []
    for p in range(1, NUM_PAGES + 1):
        base = 1 + (p - 1) * PAGE_ROWS
        slot_start = 1 + (p - 1) * DATA_ROWS_PER_PAGE
        build_page(ws, page_num=p, base_row=base, slot_start=slot_start)
        if p < NUM_PAGES:
            breaks.append(Break(id=base + PAGE_ROWS - 1, man=True))
    ws.row_breaks = RowBreak(brk=breaks)

    last_row = NUM_PAGES * PAGE_ROWS
    ws.print_area = f"A1:K{last_row}"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
