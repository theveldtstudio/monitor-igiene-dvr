import openpyxl

TEMPLATE_PATH = "public/templates/owas.xlsx"

wb = openpyxl.load_workbook(TEMPLATE_PATH)
ws = wb.active
print(f"Sheet: {ws.title}")

print(f"\n=== a) Dimensions ===")
print(f"  max_col={ws.max_column}  max_row={ws.max_row}")

print("\n=== b) Merge cells ===")
merges = sorted(str(mr) for mr in ws.merged_cells.ranges)
for m in merges:
    print(f"  {m}")
print(f"  Total: {len(merges)}")

print("\n=== c) Column widths ===")
for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"]:
    dim = ws.column_dimensions.get(col)
    w = dim.width if dim else "(not set)"
    print(f"  {col}: {w}")

print("\n=== d) Cell contents rows 1-16 ===")
for row in range(1, 17):
    for col in range(1, 11):
        cell = ws.cell(row=row, column=col)
        if cell.value is not None:
            from openpyxl.utils import get_column_letter
            cref = f"{get_column_letter(col)}{row}"
            print(f"  {cref} = {repr(cell.value)}")
