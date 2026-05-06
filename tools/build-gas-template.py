"""
Build public/templates/gas.xlsx — Foglio di campagna ufficiale Gas.
Pattern: 4 pagine x 8 misure = 32 max, identico a rumore.xlsx.

Colonne (A..N):
  A REG | [n°]
  B FASE LAVORATIVA
  C POSTAZIONE DI MISURA
  D TEMPO DI PRELIEVO
  E TIPO DI PRELIEVO
  F MACCHINE/IMPIANTI PRESENTI
  G NO2  (CONCENTRAZIONE)
  H NO
  I CO
  J CO2
  K H2S
  L Altro    (gas custom — sostituisce il NO2 duplicato del sorgente)
  M O2
  N NOTE / EVENTUALI INTERFERENZE
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.page import PageMargins

OUT = Path(__file__).resolve().parent.parent / "public" / "templates" / "gas.xlsx"

# Costanti layout
COL_WIDTHS = {
    "A": 8.5,
    "B": 18.0,
    "C": 18.0,
    "D": 12.0,
    "E": 12.0,
    "F": 22.0,
    "G": 9.0,
    "H": 9.0,
    "I": 9.0,
    "J": 9.0,
    "K": 9.0,
    "L": 9.0,
    "M": 9.0,
    "N": 26.0,
}
ROW_H_HEADER = 30.0
ROW_H_TITLE = 24.95
ROW_H_COLHDR = 22.0
ROW_H_SUBHDR = 18.0
ROW_H_DATA = 54.95
ROW_H_FOOTER = 27.75
ROW_H_EMPTY = 15.0

PAGE_ROWS = 15
DATA_ROWS_PER_PAGE = 8

THIN = Side(style="thin", color="000000")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
FONT_HDR = Font(name="Calibri", size=11, bold=True)
FONT_TITLE = Font(name="Calibri", size=12, bold=True)
FONT_NORMAL = Font(name="Calibri", size=11)
FILL_TITLE = PatternFill(fill_type="solid", fgColor="DDDDDD")
FILL_COLHDR = PatternFill(fill_type="solid", fgColor="EEEEEE")


def build_page(ws, page_num: int, base_row: int, slot_start: int):
    """Disegna una pagina del foglio campagna. base_row = riga 1-based di partenza."""
    # Riga base_row → header anagrafica
    r = base_row
    ws.row_dimensions[r].height = ROW_H_HEADER
    ws.cell(r, 1, "Impresa:").font = FONT_HDR
    # B:C committente (merge)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(r, 2).alignment = CENTER
    ws.cell(r, 4, "Data:").font = FONT_HDR
    ws.cell(r, 5, "Cantiere:").font = FONT_HDR
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=10)
    ws.cell(r, 6).alignment = CENTER
    ws.cell(r, 11, "Progressiva:").font = FONT_HDR
    ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=13)
    ws.cell(r, 14, f"Pag: {page_num}").font = FONT_HDR
    ws.cell(r, 14).alignment = Alignment(horizontal="right", vertical="center")

    # Riga +1 → titolo
    r2 = base_row + 1
    ws.row_dimensions[r2].height = ROW_H_TITLE
    ws.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=14)
    cell = ws.cell(r2, 1, "CAMPIONAMENTO GAS (Titolo IX D.Lgs 81/08)")
    cell.font = FONT_TITLE
    cell.alignment = CENTER
    cell.fill = FILL_TITLE

    # Riga +2 → header colonne (parte 1)
    r3 = base_row + 2
    ws.row_dimensions[r3].height = ROW_H_COLHDR
    headers_top = {
        1: ("REG | [n°]", 2),       # A: span su 2 righe
        2: ("FASE LAVORATIVA", 2),  # B
        3: ("POSTAZIONE DI MISURA", 2),
        4: ("TEMPO DI PRELIEVO", 2),
        5: ("TIPO DI PRELIEVO", 2),
        6: ("MACCHINE/IMPIANTI PRESENTI", 2),
        14: ("NOTE / EVENTUALI INTERFERENZE", 2),
    }
    for col, (txt, span) in headers_top.items():
        if span == 2:
            ws.merge_cells(start_row=r3, start_column=col, end_row=r3 + 1, end_column=col)
        c = ws.cell(r3, col, txt)
        c.font = FONT_HDR
        c.alignment = CENTER
        c.fill = FILL_COLHDR
    # Header CONCENTRAZIONE (G:M merge orizzontale su r3)
    ws.merge_cells(start_row=r3, start_column=7, end_row=r3, end_column=13)
    c = ws.cell(r3, 7, "CONCENTRAZIONE")
    c.font = FONT_HDR
    c.alignment = CENTER
    c.fill = FILL_COLHDR

    # Riga +3 → sub-header gas
    r4 = base_row + 3
    ws.row_dimensions[r4].height = ROW_H_SUBHDR
    subs = ["NO2", "NO", "CO", "CO2", "H2S", "Altro", "O2"]
    for i, label in enumerate(subs):
        c = ws.cell(r4, 7 + i, label)
        c.font = FONT_HDR
        c.alignment = CENTER
        c.fill = FILL_COLHDR

    # Bordi su tutta l'area header r3:r4 colonne A..N
    for rr in (r3, r4):
        for col in range(1, 15):
            ws.cell(rr, col).border = BORDER_ALL

    # Righe dati r5..r12 (8 misure)
    for i in range(DATA_ROWS_PER_PAGE):
        rr = base_row + 4 + i
        ws.row_dimensions[rr].height = ROW_H_DATA
        for col in range(1, 15):
            ws.cell(rr, col).border = BORDER_ALL
            ws.cell(rr, col).alignment = CENTER
            ws.cell(rr, col).font = FONT_NORMAL
        # Numero misura in A
        ws.cell(rr, 1, slot_start + i).font = FONT_HDR

    # Footer "Tecnico rilevatore" alla riga base_row+12
    r_foot = base_row + 12
    ws.row_dimensions[r_foot].height = ROW_H_FOOTER
    ws.merge_cells(start_row=r_foot, start_column=1, end_row=r_foot, end_column=2)
    c = ws.cell(r_foot, 1, "Tecnico rilevatore")
    c.font = FONT_HDR
    c.alignment = LEFT
    ws.merge_cells(start_row=r_foot, start_column=3, end_row=r_foot, end_column=8)
    c = ws.cell(r_foot, 3)
    c.alignment = LEFT
    ws.merge_cells(start_row=r_foot, start_column=9, end_row=r_foot, end_column=14)
    c = ws.cell(r_foot, 9, "Strumentazione:")
    c.font = FONT_NORMAL
    c.alignment = LEFT

    # Righe vuote r_foot+1, +2 separator
    for off in (13, 14):
        rr = base_row + off
        if rr <= base_row + PAGE_ROWS - 1:
            ws.row_dimensions[rr].height = ROW_H_EMPTY


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Foglio1"

    # Larghezze colonne
    for col_letter, w in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = w

    # Page setup landscape, fit to width 1, scale
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True

    # 4 pagine
    breaks = []
    for p in range(1, 5):
        base = 1 + (p - 1) * PAGE_ROWS
        slot_start = 1 + (p - 1) * DATA_ROWS_PER_PAGE
        build_page(ws, page_num=p, base_row=base, slot_start=slot_start)
        if p < 4:
            # Page break dopo l'ultima riga della pagina (base + 14, 0-based row break id)
            breaks.append(Break(id=base + PAGE_ROWS - 1, man=True))
    ws.row_breaks = RowBreak(brk=breaks)

    # Print area
    last_row = 4 * PAGE_ROWS
    ws.print_area = f"A1:N{last_row}"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"OK → {OUT}")


if __name__ == "__main__":
    main()
