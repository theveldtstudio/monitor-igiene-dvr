import openpyxl
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('public/templates/owas.xlsx')
ws = wb.active

print("=== MERGE RANGES ===")
for mr in ws.merged_cells.ranges:
    print(f"  {mr}")

print("\n=== CELLA B1 ===")
print(f"  Valore: {ws['B1'].value!r}")
print(f"  Tipo: {type(ws['B1']).__name__}")
print(f"  In merge: {any('B1' in str(mr) for mr in ws.merged_cells.ranges)}")

print("\n=== HEADER RIGA 1 ===")
for col_letter in ['A','B','C','D','E','F','G','H','I','J']:
    cell = ws[f'{col_letter}1']
    print(f"  {col_letter}1: value={cell.value!r}, numFmt={cell.number_format!r}")

print("\n=== RIGA 5 (prima riga dati) — numFmt ===")
for col_letter in ['A','B','C','D','E','F','G','H','I','J']:
    cell = ws[f'{col_letter}5']
    print(f"  {col_letter}5: numFmt={cell.number_format!r}, value={cell.value!r}")

print("\n=== RIGHE DATI 5-12: numFmt colonna D vs E-I ===")
for r in range(5, 13):
    d_fmt = ws[f'D{r}'].number_format
    e_fmt = ws[f'E{r}'].number_format
    f_fmt = ws[f'F{r}'].number_format
    print(f"  Riga {r}: D={d_fmt!r}, E={e_fmt!r}, F={f_fmt!r}")

print("\n=== CELL STYLES (xfId proxy via _style) riga 5 ===")
for col_letter in ['A','B','C','D','E','F','G','H','I','J']:
    cell = ws[f'{col_letter}5']
    style_id = getattr(cell, '_style', None)
    xf = getattr(style_id, 'xfId', 'N/A') if style_id else 'N/A'
    print(f"  {col_letter}5: _style={style_id!r}")

wb.close()
